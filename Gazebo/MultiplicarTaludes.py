#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MultiplicarTaludes.py - Script para multiplicar largura e altura dos taludes por 2
"""

import os
import pandas as pd
from openpyxl import load_workbook


class MultiplicarTaludes:
    def __init__(self):
        # Obter o diretório do script
        self.script_dir = os.path.dirname(os.path.abspath(__file__))
        
        # Caminhos dos arquivos
        self.todos_csv = os.path.join(self.script_dir, 'todos_pontos_gps.csv')
        self.todos_xlsx = os.path.join(self.script_dir, 'todos_pontos_gps.xlsx')
        
    def multiplicar_csv(self):
        """Multiplica a largura e altura dos taludes no CSV por 2"""
        print("\n📝 Processando arquivo CSV...")
        
        try:
            # Ler o arquivo
            df = pd.read_csv(self.todos_csv)
            
            print(f"  📊 Total de registros: {len(df)}")
            
            # Encontrar taludes
            taludes_mask = df['ID'].str.contains('talude', case=False, na=False)
            num_taludes = taludes_mask.sum()
            print(f"  🎯 Taludes encontrados: {num_taludes}")
            
            # Multiplicar Vx_largura e Vy_altura dos taludes por 2
            df.loc[taludes_mask, 'Vx_largura'] = df.loc[taludes_mask, 'Vx_largura'] * 2
            df.loc[taludes_mask, 'Vy_altura'] = df.loc[taludes_mask, 'Vy_altura'] * 2
            
            # Salvar o arquivo
            df.to_csv(self.todos_csv, index=False)
            print(f"  ✅ Arquivo CSV atualizado com sucesso!")
            
            return True
            
        except Exception as e:
            print(f"  ❌ Erro ao processar CSV: {e}")
            return False
    
    def multiplicar_xlsx(self):
        """Multiplica a largura e altura dos taludes no XLSX por 2"""
        print("\n📊 Processando arquivo XLSX...")
        
        try:
            # Ler o arquivo
            df = pd.read_excel(self.todos_xlsx, sheet_name='Sheet1', engine='openpyxl')
            
            print(f"  📊 Total de registros: {len(df)}")
            
            # Encontrar taludes
            taludes_mask = df['ID'].str.contains('talude', case=False, na=False)
            num_taludes = taludes_mask.sum()
            print(f"  🎯 Taludes encontrados: {num_taludes}")
            
            # Multiplicar Vx_largura e Vy_altura dos taludes por 2
            df.loc[taludes_mask, 'Vx_largura'] = df.loc[taludes_mask, 'Vx_largura'] * 2
            df.loc[taludes_mask, 'Vy_altura'] = df.loc[taludes_mask, 'Vy_altura'] * 2
            
            # Carregar o workbook
            wb = load_workbook(self.todos_xlsx)
            
            # Remover a aba Sheet1 existente
            if 'Sheet1' in wb.sheetnames:
                del wb['Sheet1']
            
            # Criar nova aba com os dados atualizados
            ws = wb.create_sheet('Sheet1', 0)
            
            # Escrever headers
            for col_idx, col_name in enumerate(df.columns, 1):
                ws.cell(row=1, column=col_idx, value=col_name)
            
            # Escrever dados
            for row_idx, row in enumerate(df.itertuples(index=False), 2):
                for col_idx, value in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            
            # Salvar o workbook
            wb.save(self.todos_xlsx)
            print(f"  ✅ Arquivo XLSX atualizado com sucesso!")
            
            return True
            
        except Exception as e:
            print(f"  ❌ Erro ao processar XLSX: {e}")
            return False
    
    def run(self):
        """Função principal"""
        print("=" * 60)
        print("🔢 MULTIPLICANDO LARGURA E ALTURA DOS TALUDES POR 2")
        print("=" * 60)
        
        try:
            # Multiplicar CSV
            csv_ok = self.multiplicar_csv()
            
            # Multiplicar XLSX
            xlsx_ok = self.multiplicar_xlsx()
            
            # Resumo
            print("\n" + "=" * 60)
            if csv_ok and xlsx_ok:
                print("✅ PROCESSO CONCLUÍDO COM SUCESSO!")
                print("=" * 60)
                print(f"📁 CSV: {self.todos_csv}")
                print(f"📁 XLSX: {self.todos_xlsx}")
                print("\n📝 Alterações:")
                print("   - Vx_largura multiplicado por 2 (apenas taludes)")
                print("   - Vy_altura multiplicado por 2 (apenas taludes)")
            else:
                print("⚠️  ERRO NO PROCESSAMENTO")
                print("=" * 60)
                
        except Exception as e:
            print(f"\n❌ ERRO: {e}")
            return False
        
        return True


def main():
    try:
        mult = MultiplicarTaludes()
        sucesso = mult.run()
        exit(0 if sucesso else 1)
    except Exception as e:
        print(f"❌ Erro fatal: {e}")
        exit(1)


if __name__ == '__main__':
    main()

